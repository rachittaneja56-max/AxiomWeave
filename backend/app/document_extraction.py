"""Bounded extraction for user supplied text, DOCX, and PDF documents."""

import base64
from collections.abc import Awaitable, Callable
from dataclasses import dataclass
from io import BytesIO
from typing import Any, cast

import pymupdf
from docx import Document
from docx.text.paragraph import Paragraph

from app.source_versions import normalize_source_text

MAX_SOURCE_CHARACTERS = 20_000
MAX_TEXT_FILE_BYTES = 80 * 1024
MAX_PDF_BYTES = 8 * 1024 * 1024
MAX_PDF_PAGES = 20
MAX_OCR_PAGES = 8
OCR_NATIVE_TEXT_THRESHOLD = 40
_PDF: Any = pymupdf

OCRProvider = Callable[[int, str], Awaitable[str]]


class DocumentExtractionError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    source_text: str
    media_type: str
    extraction_method: str
    page_count: int | None = None
    ocr_used: bool = False


def _normalize(text: str) -> str:
    try:
        return normalize_source_text(text)
    except ValueError as error:
        if "non-whitespace" in str(error):
            raise DocumentExtractionError(
                "no_readable_text", "The document contains no readable text."
            ) from None
        raise DocumentExtractionError(
            "source_too_long", "Extracted text exceeds the 20,000-character source limit."
        ) from None


def extract_docx(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
        parts: list[str] = []
        items = list(document.iter_inner_content())
        for item in items:
            if isinstance(item, Paragraph):
                parts.append(item.text)
            else:
                for row in item.rows:
                    parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception:
        raise DocumentExtractionError(
            "invalid_document", "The DOCX file could not be read."
        ) from None


def _read_pdf(content: bytes) -> tuple[pymupdf.Document, list[str]]:
    if len(content) > MAX_PDF_BYTES:
        raise DocumentExtractionError("file_too_large", "PDF files must be 8 MiB or smaller.")
    try:
        pdf = _PDF.open(stream=content, filetype="pdf")
    except Exception:
        raise DocumentExtractionError("invalid_pdf", "The PDF file could not be read.") from None
    if pdf.is_encrypted:
        pdf.close()
        raise DocumentExtractionError("encrypted_pdf", "Password-protected PDFs are not supported.")
    if len(pdf) == 0:
        pdf.close()
        raise DocumentExtractionError("empty_pdf", "The PDF does not contain any pages.")
    if len(pdf) > MAX_PDF_PAGES:
        pdf.close()
        raise DocumentExtractionError("too_many_pages", "PDF files may contain up to 20 pages.")
    return pdf, [cast(str, page.get_text("text", sort=True)) for page in pdf]


async def extract_pdf(content: bytes, ocr_provider: OCRProvider | None) -> ExtractedDocument:
    pdf, native_pages = _read_pdf(content)
    try:
        needs_ocr = [
            index
            for index, text in enumerate(native_pages)
            if len("".join(text.split())) < OCR_NATIVE_TEXT_THRESHOLD
        ]
        if len(needs_ocr) > MAX_OCR_PAGES:
            raise DocumentExtractionError(
                "too_many_ocr_pages", "OCR is limited to 8 scanned pages per PDF."
            )
        if needs_ocr and ocr_provider is None:
            raise DocumentExtractionError(
                "ocr_not_configured",
                "This PDF contains scanned pages. Scanned page reading is not configured.",
            )
        ocr_used = False
        for index in needs_ocr:
            page: Any = pdf[index]
            scale = min(2.0, 2400 / max(page.rect.width, page.rect.height))
            pixmap = page.get_pixmap(matrix=_PDF.Matrix(scale, scale), alpha=False)
            image_data = base64.b64encode(pixmap.tobytes("png")).decode("ascii")
            assert ocr_provider is not None
            try:
                native_pages[index] = await ocr_provider(index + 1, image_data)
            except Exception:
                raise DocumentExtractionError(
                    "ocr_failed", "Scanned pages could not be read. Please try another PDF."
                ) from None
            ocr_used = True
        if not any(text.strip() for text in native_pages):
            raise DocumentExtractionError(
                "no_readable_text", "No readable text was found in the PDF."
            )
        marked = "\n\n".join(
            f"# Page {number}\n\n{text}" for number, text in enumerate(native_pages, 1)
        )
        source_text = _normalize(marked)
        return ExtractedDocument(
            source_text=source_text,
            media_type="application/pdf",
            extraction_method="pdf_native_plus_ocr" if ocr_used else "pdf_native",
            page_count=len(native_pages),
            ocr_used=ocr_used,
        )
    finally:
        pdf.close()


async def extract_document(
    filename: str, content: bytes, media_type: str, ocr_provider: OCRProvider | None = None
) -> ExtractedDocument:
    extension = filename.rpartition(".")[2].lower()
    if extension in {"txt", "md"}:
        if len(content) > MAX_TEXT_FILE_BYTES:
            raise DocumentExtractionError("file_too_large", "Text files must be 80 KiB or smaller.")
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise DocumentExtractionError(
                "invalid_encoding", "The file must contain valid UTF-8 text."
            ) from None
        kind = "text/markdown" if extension == "md" else "text/plain"
        return ExtractedDocument(_normalize(text), kind, "text")
    if extension == "docx":
        if len(content) > MAX_PDF_BYTES:
            raise DocumentExtractionError("file_too_large", "DOCX files must be 8 MiB or smaller.")
        return ExtractedDocument(
            _normalize(extract_docx(content)),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        )
    if extension == "pdf":
        return await extract_pdf(content, ocr_provider)
    raise DocumentExtractionError("unsupported_file", "Upload a TXT, MD, DOCX, or PDF file.")

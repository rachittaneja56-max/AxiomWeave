from io import BytesIO
from typing import Any

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.api.sources import MAX_TEXT_FILE_BYTES
from app.domain.transformation import SOURCE_TEXT_MAX_LENGTH
from app.main import app

client: TestClient = TestClient(app)
PDF: Any = pymupdf


@pytest.fixture(autouse=True)
def use_authenticated_test_client(authorized_client: TestClient) -> None:
    global client
    client = authorized_client


def upload(filename: str, content: bytes, content_type: str = "application/octet-stream"):
    return client.post(
        "/api/sources/text-file",
        files={"file": (filename, content, content_type)},
    )


def upload_document(filename: str, content: bytes, content_type: str):
    return client.post("/api/sources/file", files={"file": (filename, content, content_type)})


def test_extracts_utf8_txt_with_canonical_metadata() -> None:
    source_text = "Café opens on Saturday. 🌱"

    response = upload("report.txt", source_text.encode("utf-8"), "text/plain")

    assert response.status_code == 200
    assert response.json() == {
        "filename": "report.txt",
        "media_type": "text/plain",
        "character_count": len(source_text),
        "source_text": source_text,
        "extraction_method": "text",
        "page_count": None,
        "ocr_used": False,
    }


def test_extracts_markdown_and_supports_utf8_bom() -> None:
    source_text = "# Project update\n\nThe team is ready."

    response = upload("update.md", b"\xef\xbb\xbf" + source_text.encode("utf-8"), "text/markdown")

    assert response.status_code == 200
    assert response.json()["source_text"] == source_text
    assert response.json()["media_type"] == "text/markdown"


def test_new_upload_endpoint_extracts_docx_paragraph_and_table() -> None:
    document = Document()
    document.add_heading("Project note", level=1)
    document.add_paragraph("The hall opens at ten.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Location"
    table.cell(0, 1).text = "Riverdale"
    output = BytesIO()
    document.save(output)

    response = upload_document(
        "brief.docx",
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert response.status_code == 200
    assert response.json()["extraction_method"] == "docx"
    assert "Project note" in response.json()["source_text"]
    assert "Location\tRiverdale" in response.json()["source_text"]


def test_new_upload_endpoint_extracts_native_pdf_with_page_locator() -> None:
    pdf = PDF.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "A native text PDF page contains this readable report content.")
    content = pdf.tobytes()
    pdf.close()

    response = upload_document("report.pdf", content, "application/pdf")

    assert response.status_code == 200
    assert response.json()["extraction_method"] == "pdf_native"
    assert response.json()["page_count"] == 1
    assert "# Page 1" in response.json()["source_text"]


def test_scanned_pdf_without_configuration_returns_specific_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf = PDF.open()
    pdf.new_page()
    content = pdf.tobytes()
    pdf.close()
    monkeypatch.setattr(
        "app.api.sources.get_settings",
        lambda: type("SettingsStub", (), {"openai_api_key": None})(),
    )

    response = upload_document("scan.pdf", content, "application/pdf")

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "ocr_not_configured"


def test_txt_extraction_uses_the_same_canonical_newlines_as_pasted_text() -> None:
    response = upload("source.txt", b"\xef\xbb\xbfFirst\r\nsecond\rthird", "text/plain")

    assert response.status_code == 200
    assert response.json()["source_text"] == "First\nsecond\nthird"
    assert response.json()["character_count"] == len(response.json()["source_text"])


@pytest.mark.parametrize(
    ("content", "code"),
    [(b"", "empty_source"), (b" \n\t", "empty_source"), (b"\xff", "invalid_encoding")],
)
def test_rejects_empty_whitespace_and_invalid_utf8(content: bytes, code: str) -> None:
    response = upload("source.txt", content)

    assert response.status_code == 422
    assert response.json()["error"]["code"] == code


def test_rejects_unsupported_extension() -> None:
    response = upload("report.pdf", b"not a PDF")

    assert response.status_code == 415
    assert response.json()["error"]["code"] == "unsupported_file"


def test_rejects_incompatible_content_type() -> None:
    response = upload("report.txt", b"not a text file", "application/pdf")

    assert response.status_code == 415
    assert response.json()["error"]["code"] == "unsupported_media_type"


def test_rejects_raw_file_over_byte_limit() -> None:
    response = upload("large.txt", b"a" * (MAX_TEXT_FILE_BYTES + 1))

    assert response.status_code == 413
    assert response.json()["error"]["code"] == "file_too_large"


def test_rejects_extracted_source_over_canonical_limit() -> None:
    response = upload("large.txt", b"a" * (SOURCE_TEXT_MAX_LENGTH + 1))

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "source_too_long"


def test_rejects_missing_file() -> None:
    response = client.post("/api/sources/text-file", files={})

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_request"

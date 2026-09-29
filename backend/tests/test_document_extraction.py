import asyncio
from io import BytesIO
from typing import Any

import pymupdf
import pytest
from docx import Document

from app.document_extraction import (
    MAX_OCR_PAGES,
    MAX_PDF_BYTES,
    DocumentExtractionError,
    extract_document,
    extract_pdf,
)

PDF: Any = pymupdf


def make_pdf(pages: list[str | None]) -> bytes:
    pdf = PDF.open()
    for text in pages:
        page = pdf.new_page(width=612, height=792)
        if text:
            page.insert_text((72, 72), text)
    return pdf.tobytes()


def make_scanned_pdf() -> bytes:
    source = PDF.open()
    page = source.new_page(width=612, height=792)
    page.insert_text((72, 72), "Scanned page content", fontsize=24)
    pixmap = page.get_pixmap(matrix=PDF.Matrix(2, 2), alpha=False)
    scanned = PDF.open()
    image_page = scanned.new_page(width=612, height=792)
    image_page.insert_image(image_page.rect, stream=pixmap.tobytes("png"))
    result = scanned.tobytes()
    source.close()
    scanned.close()
    return result


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Project update", level=1)
    document.add_paragraph("Paragraph content follows.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Budget"
    table.cell(0, 1).text = "Approved"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_extracts_paragraphs_heading_and_table_cells() -> None:
    result = asyncio.run(extract_document("brief.docx", make_docx(), "application/octet-stream"))
    assert result.extraction_method == "docx"
    assert "Project update" in result.source_text
    assert "Paragraph content follows." in result.source_text
    assert "Budget\tApproved" in result.source_text


def test_native_pdf_keeps_page_markers_and_skips_ocr() -> None:
    async def fail_ocr(_page: int, _image: str) -> str:
        raise AssertionError("native PDF must not call OCR")

    result = asyncio.run(
        extract_document(
            "native.pdf",
            make_pdf(
                [
                    "A sufficiently long native text page for extraction.",
                    "The second native page contains enough useful text too.",
                ]
            ),
            "application/pdf",
            fail_ocr,
        )
    )
    assert result.extraction_method == "pdf_native"
    assert result.page_count == 2
    assert "# Page 1" in result.source_text
    assert "# Page 2" in result.source_text
    assert result.source_text.index("# Page 1") < result.source_text.index("# Page 2")
    assert result.ocr_used is False


def test_mixed_pdf_ocr_only_scanned_page_and_merges_in_order() -> None:
    async def fake_ocr(page_number: int, _image: str) -> str:
        assert page_number == 2
        return "Scanned page transcribed."

    result = asyncio.run(
        extract_pdf(
            make_pdf(["Native page text with enough characters to avoid any OCR.", None]), fake_ocr
        )
    )
    assert result.extraction_method == "pdf_native_plus_ocr"
    assert result.ocr_used is True
    assert result.source_text.index("Native page text") < result.source_text.index(
        "Scanned page transcribed"
    )


def test_scanned_pdf_without_provider_has_safe_configuration_error() -> None:
    with pytest.raises(DocumentExtractionError) as error:
        asyncio.run(extract_pdf(make_scanned_pdf(), None))
    assert error.value.code == "ocr_not_configured"


def test_ocr_provider_failure_is_safe() -> None:
    async def broken_ocr(_page: int, _image: str) -> str:
        raise RuntimeError("private provider detail")

    with pytest.raises(DocumentExtractionError) as error:
        asyncio.run(extract_pdf(make_scanned_pdf(), broken_ocr))
    assert error.value.code == "ocr_failed"
    assert "private provider detail" not in str(error.value)


def test_empty_ocr_result_is_not_saved_as_page_marker_only_source() -> None:
    async def empty_ocr(_page: int, _image: str) -> str:
        return " \n "

    with pytest.raises(DocumentExtractionError) as error:
        asyncio.run(extract_pdf(make_scanned_pdf(), empty_ocr))
    assert error.value.code == "no_readable_text"


@pytest.mark.parametrize("content", [b"not pdf", b""])
def test_corrupt_or_empty_pdf_is_rejected(content: bytes) -> None:
    with pytest.raises(DocumentExtractionError):
        asyncio.run(extract_pdf(content, None))


def test_encrypted_pdf_is_rejected() -> None:
    pdf = PDF.open()
    pdf.new_page()
    encrypted = pdf.tobytes(encryption=PDF.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret")
    pdf.close()
    with pytest.raises(DocumentExtractionError) as error:
        asyncio.run(extract_pdf(encrypted, None))
    assert error.value.code == "encrypted_pdf"


def test_pdf_page_and_ocr_page_limits_are_explicit() -> None:
    pdf = PDF.open()
    for _ in range(21):
        pdf.new_page()
    too_many = pdf.tobytes()
    pdf.close()
    with pytest.raises(DocumentExtractionError) as error:
        asyncio.run(extract_pdf(too_many, None))
    assert error.value.code == "too_many_pages"
    assert MAX_OCR_PAGES == 8


def test_rejects_pdf_over_raw_file_size_limit() -> None:
    with pytest.raises(DocumentExtractionError) as error:
        asyncio.run(extract_pdf(b"x" * (MAX_PDF_BYTES + 1), None))
    assert error.value.code == "file_too_large"


def test_normalized_source_limit_rejects_without_truncation() -> None:
    with pytest.raises(DocumentExtractionError) as error:
        asyncio.run(extract_document("long.txt", b"x" * 20_001, "text/plain"))
    assert error.value.code == "source_too_long"
